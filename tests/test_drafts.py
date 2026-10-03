"""Letter and notice drafts, the Rule 80K packet and .docx export (app/routers/drafts.py)."""

from __future__ import annotations

import io
import re

import pytest
from docx import Document

from app import config, llm
from app.routers import drafts as d
from tests.conftest import CSRF

CASE = {
    "address": "12 Main Street",
    "map_lot": "041-112",
    "owner": "Main Place LLC",
    "tags": ["property-maintenance"],
    "status": "open",
    "summary": "Junk and two unregistered vehicles in the side yard.",
}

NOV_VALUES = {
    "owner_name": "Main Place LLC",
    "property_address": "12 Main Street",
    "owner_address": "PO Box 1\nWaterville, ME 04901",
    "violation_type": "zoning_specific",
    "inspection_date": "2026-09-30",
    "section_cited": "§ 275-4.26B(2)",
    "facts": "A fence about 8 feet high stands along the rear lot line.",
    "corrective_action": "Lower the fence to 6 feet or less, or apply for a permit.",
    "correct_by": "2026-10-31",
    "letter_date": "2026-10-03",
    "signer_name": "Adam Example",
}


def make_case(c, **over) -> dict:
    r = c.post("/api/staff/cases", json={**CASE, **over}, headers=CSRF)
    assert r.status_code == 201, r.text
    return r.json()


def make_draft(c, template="nov-1", values=None, **over) -> dict:
    r = c.post("/api/staff/drafts", json={"template": template, "values": values or {}, **over}, headers=CSRF)
    assert r.status_code == 201, r.text
    return r.json()


def body_of(tid: str, values: dict) -> d.Rendered:
    t = d.templates()[tid]
    vals = {**d.default_values(t), **values}
    return d.render(t, d.check_values(t, vals))


# ---------------------------------------------------------------- templates


EXPECTED = {"nov-1", "nov-2", "nov-3", "stop-work", "abutter-se", "abutter-se-decision", "zba-hearing", "decision",
            "80k-packet", "80k-memo"}


def test_every_template_loads_and_uses_only_known_placeholders():
    ts = d.templates()
    assert EXPECTED <= set(ts)
    for tid, t in ts.items():
        assert t.title and t.description and t.notes, tid
        assert not (d.placeholders_in(t) - d.known_names(t)), (tid, d.placeholders_in(t) - d.known_names(t))
        clash = {k for opts in t.variants.values() for texts in opts.values() for k in texts} & set(t.field_map)
        assert not clash, (tid, clash)
        for name in ("letter_date", "signer_name", "signer_title"):
            assert name in t.field_map, (tid, name)
        # Defaults alone render with no stray {placeholder}.
        r = d.render(t, d.default_values(t))
        assert not d.PLACEHOLDER_RE.search(r.body), (tid, d.PLACEHOLDER_RE.findall(r.body))


def test_no_em_dashes_or_emoji_in_templates():
    for p in d.TEMPLATE_DIR.iterdir():
        text = p.read_text(encoding="utf-8")
        assert "—" not in text, p.name
        assert not re.search(r"[\U0001F300-\U0001FAFF☀-⛿✀-➿]", text), p.name


def test_enforcement_templates_carry_appeal_penalty_and_deadline():
    for tid in ("nov-1", "nov-2", "nov-3", "stop-work"):
        t = d.templates()[tid]
        assert "violation_type" in t.field_map
        assert "{appeal_text}" in t.body and "{penalty_text}" in t.body, tid


# ---------------------------------------------------------------- rendering


def test_zoning_nov_cites_zba_appeal_and_penalty_tier():
    r = body_of("nov-1", NOV_VALUES)
    assert not r.missing
    b = r.body
    assert "§ 275-4.26B(2)" in b
    assert "§ 275-6.2E(1)" in b and "within 30 days" in b
    assert "November 2, 2026" in b  # 30 days after October 3
    assert "30-A M.R.S. § 4452(3)(B)" in b and "$5,000" in b
    assert "30-A M.R.S. § 2691(4)" in b
    assert "§ 275-6.1A(3)(a)" in b  # Solicitor and Council copy
    assert "October 31, 2026" in b
    assert "Mortgagee" not in b
    assert "[VERIFY" not in b  # every zoning statement here was checked


def test_property_maintenance_nov_needs_mortgagee_and_uses_205_penalty():
    r = body_of("nov-1", {**NOV_VALUES, "violation_type": "property_maintenance", "section_cited": "§ 205-4C"})
    assert {m["name"] for m in r.missing} == {"mortgagee_name", "mortgagee_address"}
    assert "[MISSING: Mortgagee]" in r.body
    r = body_of("nov-1", {**NOV_VALUES, "violation_type": "property_maintenance", "section_cited": "§ 205-4C",
                          "mortgagee_name": "First Bank", "mortgagee_address": "1 Bank St"})
    assert not r.missing
    assert "§ 205-7" in r.body and "§ 205-8" in r.body and "$100" in r.body
    assert "cc: Mortgagee: First Bank, 1 Bank St (§ 205-7)" in r.body
    assert "[VERIFY" in r.body  # the Ch. 205 appeal route is not settled


def test_building_no_permit_nov_has_tier_a_and_after_the_fact_fee():
    r = body_of("nov-1", {**NOV_VALUES, "violation_type": "building_no_permit", "section_cited": "§ 127-2A"})
    assert "§ 4452(3)(A)" in r.body and "$2,500" in r.body
    assert "§ 127-3B" in r.body and "$1,000" in r.body


def test_resource_protection_flags_conflicting_maximums():
    r = body_of("nov-1", {**NOV_VALUES, "violation_type": "zoning_rp"})
    assert "$10,000" in r.body and "$5,000" in r.body and "[VERIFY" in r.body


def test_prior_conviction_adds_paragraph_f():
    assert "$25,000" not in body_of("nov-1", NOV_VALUES).body
    assert "§ 4452(3)(F)" in body_of("nov-1", {**NOV_VALUES, "prior_conviction": "yes"}).body


def test_missing_fields_are_marked_and_listed():
    t = d.templates()["nov-1"]
    r = d.render(t, {})
    names = {m["name"] for m in r.missing}
    assert {"owner_name", "owner_address", "violation_type", "facts", "correct_by", "signer_name", "letter_date"} <= names
    assert "[MISSING: Facts observed]" in r.body
    assert "[MISSING: Ordinance and violation type]" in r.body  # variant text stands in for the missing select
    assert "{" not in r.body


def test_optional_lines_drop_and_else_lines_appear():
    r = body_of("nov-1", NOV_VALUES)
    assert "City case:" not in r.body and "in person" not in r.body
    r = body_of("nov-1", {**NOV_VALUES, "oral_notice_date": "2026-09-29", "case_ref": "2026-abcdef"})
    assert "September 29, 2026" in r.body and "City case: 2026-abcdef" in r.body
    # Decision: blank conclusions give the § 275-5.20B skeleton, typed ones replace it.
    base = {"application_kind": "special_exception", "decision": "approved"}
    assert "§ 275-5.20B(3)" in body_of("decision", base).body
    typed = body_of("decision", {**base, "conclusions": "All eight standards are met."}).body
    assert "All eight standards are met." in typed and "§ 275-5.20B(3)" not in typed


def test_required_when_follows_other_fields():
    assert "conditions" not in {m["name"] for m in body_of("decision", {"decision": "approved"}).missing}
    assert "conditions" in {m["name"] for m in body_of("decision", {"decision": "approved_conditions"}).missing}


def test_checks_compare_dates():
    r = body_of("zba-hearing", {"filed_date": "2026-09-01", "hearing_date": "2026-10-30", "letter_date": "2026-10-01"})
    by = {c["message"][:20]: c["ok"] for c in r.checks}
    assert by["The hearing is withi"] is False  # 59 days after filing
    assert by["This notice is dated"] is True
    r = body_of("abutter-se", {"mailing_date": "2026-10-01"})
    assert "October 15, 2026" in r.body


def test_check_values_rejects_unknown_fields_and_bad_input():
    t = d.templates()["nov-1"]
    with pytest.raises(Exception) as e:
        d.check_values(t, {"nope": "x"})
    assert e.value.status_code == 422
    with pytest.raises(Exception) as e:
        d.check_values(t, {"correct_by": "next week"})
    assert e.value.status_code == 422
    with pytest.raises(Exception) as e:
        d.check_values(t, {"violation_type": "made_up"})
    assert e.value.status_code == 422


# ---------------------------------------------------------------- penalty helper


def test_penalty_range_inclusive_days():
    r = d.penalty_range("B", "2026-08-01", "2026-08-31", "150")
    assert r["days"] == 31
    assert (r["min_total"], r["max_total"], r["requested_total"]) == (3100, 155000, 4650)
    assert not r["warnings"]
    r = d.penalty_range("A", "2026-08-01", "2026-08-01", "3000")
    assert r["days"] == 1 and r["max_total"] == 2500
    assert r["warnings"] and "outside" in r["warnings"][0]
    r = d.penalty_range("205", "2026-08-10", "2026-08-01")
    assert r["days"] is None and r["warnings"]
    assert d.penalty_range("B-1", "2026-01-01", "2026-01-02")["verify"].startswith("[VERIFY")
    with pytest.raises(ValueError):
        d.penalty_range("Z", None, None)


def test_penalty_endpoints(staff_client):
    tiers = staff_client.get("/api/staff/drafts/penalty-tiers").json()["tiers"]
    assert [t["value"] for t in tiers] == ["A", "B", "B-1", "F", "205"]
    r = staff_client.post("/api/staff/drafts/penalty", json={"tier": "205", "start": "2026-09-01", "end": "2026-09-10"}, headers=CSRF)
    assert r.status_code == 200 and r.json()["min_total"] == 1000
    assert staff_client.post("/api/staff/drafts/penalty", json={"tier": "Q"}, headers=CSRF).status_code == 422
    assert staff_client.post("/api/staff/drafts/penalty", json={"tier": "A", "start": "x"}, headers=CSRF).status_code == 422


def test_packet_computes_hearing_and_penalty():
    r = body_of("80k-packet", {"violation_type": "zoning_specific", "penalty_tier": "B", "violation_start": "2026-08-01",
                               "violation_end": "2026-09-30", "service_date": "2026-10-05", "hearing_date": "2026-10-20"})
    assert "61 days" in r.body and "$6,100 to $305,000" in r.body
    # Rule 80K(b)(3): "as soon as practicable after service"; no computed 20-day dates.
    assert "as soon as practicable after service (M.R. Civ. P. 80K(b)(3))" in r.body
    assert "earliest hearing date" not in r.body and "within 20 days of service" not in r.body
    assert "by any appropriate method provided in Rule 4 (M.R. Civ. P. 80K(b)(2))" in r.body
    assert "in hand" not in r.body
    assert "Registry of Deeds" in r.body and "attested by the City Clerk" in r.body
    assert not any("20 days" in c["message"] for c in r.checks)


# ---------------------------------------------------------------- HTML


def test_html_escapes_and_marks_flags():
    html = d.to_html("# Title\n\nA <script>alert(1)</script> **bold** [VERIFY: x] [MISSING: Y]\nline two\n\n- [ ] item\n- [x] done\n\n1. one\n2. two\n\n> quote")
    assert "<script>" not in html and "&lt;script&gt;" in html
    assert "<strong>bold</strong>" in html
    assert '<mark class="flag verify">[VERIFY: x]</mark>' in html
    assert '<mark class="flag missing">[MISSING: Y]</mark>' in html
    assert "<br>line two" in html
    assert 'class="checklist"' in html and "☐" in html and "☒" in html
    assert "<ol><li>one</li><li>two</li></ol>" in html
    assert "<blockquote>" in html and "<h2>Title</h2>" in html


# ---------------------------------------------------------------- API


ENDPOINTS = [
    ("GET", "/api/staff/drafts"),
    ("POST", "/api/staff/drafts"),
    ("GET", "/api/staff/drafts/templates"),
    ("GET", "/api/staff/drafts/templates/nov-1"),
    ("POST", "/api/staff/drafts/render"),
    ("GET", "/api/staff/drafts/penalty-tiers"),
    ("POST", "/api/staff/drafts/penalty"),
    ("POST", "/api/staff/drafts/facts"),
    ("GET", "/api/staff/drafts/d2026-abcdef"),
    ("PUT", "/api/staff/drafts/d2026-abcdef"),
    ("DELETE", "/api/staff/drafts/d2026-abcdef"),
    ("GET", "/api/staff/drafts/d2026-abcdef/export.docx"),
]


@pytest.mark.parametrize("method,path", ENDPOINTS)
def test_needs_session(client, method, path):
    r = client.request(method, path, json={} if method in ("POST", "PUT") else None, headers=CSRF)
    assert r.status_code == 401


@pytest.mark.parametrize("method,path", [e for e in ENDPOINTS if e[0] != "GET"])
def test_writes_need_csrf_header(staff_client, method, path):
    r = staff_client.request(method, path, json={"template": "nov-1"} if method in ("POST", "PUT") else None)
    assert r.status_code == 403


def test_template_list_and_case_defaults(staff_client):
    r = staff_client.get("/api/staff/drafts/templates").json()
    assert r["banner"] == d.BANNER
    ids = [t["id"] for t in r["templates"]]
    assert ids[0] == "nov-1" and "80k-packet" in ids
    packet = next(t for t in r["templates"] if t["id"] == "80k-packet")
    assert packet["tools"] == ["penalty"] and packet["letterhead"] is False
    case = make_case(staff_client)
    t = staff_client.get(f"/api/staff/drafts/templates/nov-1?case_id={case['id']}").json()
    assert t["defaults"]["property_address"] == "12 Main Street"
    assert t["defaults"]["owner_name"] == "Main Place LLC"
    assert t["defaults"]["case_ref"] == case["id"]
    assert t["defaults"]["office_phone"] == "207-680-4208"
    assert staff_client.get("/api/staff/drafts/templates/nope").status_code == 404
    assert staff_client.get("/api/staff/drafts/templates/nov-1?case_id=2026-000000").status_code == 404


def test_render_endpoint(staff_client):
    r = staff_client.post("/api/staff/drafts/render", json={"template": "nov-1", "values": NOV_VALUES}, headers=CSRF)
    assert r.status_code == 200, r.text
    j = r.json()
    # Only the values sent count: defaults such as the delivery method are the client's job.
    assert {"delivery_method", "signer_title"} <= {m["name"] for m in j["missing"]}
    assert "owner_name" not in {m["name"] for m in j["missing"]}
    assert "draft-letterhead" in j["html"] and "7 College Avenue" in j["html"]
    assert j["derived"]["appeal_by"] == "2026-11-02"
    # A body sent back is previewed as is.
    j = staff_client.post("/api/staff/drafts/render", json={"template": "nov-1", "values": {}, "body": "Hello **there**"}, headers=CSRF).json()
    assert "<strong>there</strong>" in j["html"]
    bad = staff_client.post("/api/staff/drafts/render", json={"template": "nov-1", "values": {"x": "1"}}, headers=CSRF)
    assert bad.status_code == 422


def test_draft_lifecycle_and_case_timeline(staff_client):
    case = make_case(staff_client)
    from_case = {k: v for k, v in NOV_VALUES.items() if k not in ("owner_name", "property_address")}
    dr = make_draft(staff_client, values=from_case, case_id=case["id"])
    assert dr["id"].startswith("d") and dr["case_id"] == case["id"]
    assert dr["values"]["property_address"] == "12 Main Street"  # filled from the case
    assert dr["title"] == "First notice of violation: 12 Main Street"
    assert dr["edited"] is False and dr["status"] == "draft"
    assert dr["banner"] == d.BANNER and not dr["missing"]

    items = staff_client.get(f"/api/staff/cases/{case['id']}").json()["items"]
    assert [(i["kind"], i["draft_id"], i["template"]) for i in items] == [("draft", dr["id"], "nov-1")]

    listed = staff_client.get(f"/api/staff/drafts?case_id={case['id']}").json()["drafts"]
    assert [x["id"] for x in listed] == [dr["id"]] and listed[0]["template_title"] == "First notice of violation"
    assert staff_client.get("/api/staff/drafts?case_id=2026-ffffff").json()["drafts"] == []

    # Field changes rewrite the text while it is unedited.
    up = staff_client.put(f"/api/staff/drafts/{dr['id']}", json={"values": {**dr["values"], "correct_by": "2026-11-15"}}, headers=CSRF).json()
    assert "November 15, 2026" in up["body"]
    # Typing in the text marks it edited; later field changes leave it alone.
    edited_body = up["body"].replace("Sincerely,", "Very truly yours,")
    up = staff_client.put(f"/api/staff/drafts/{dr['id']}", json={"body": edited_body}, headers=CSRF).json()
    assert up["edited"] is True and "Very truly yours," in up["body"]
    up = staff_client.put(f"/api/staff/drafts/{dr['id']}", json={"values": {**up["values"], "correct_by": "2026-12-01"}}, headers=CSRF).json()
    assert "December 1, 2026" not in up["body"] and up["values"]["correct_by"] == "2026-12-01"
    # Rebuild from fields.
    rebuilt = d.render(d.templates()["nov-1"], up["values"]).body
    up = staff_client.put(f"/api/staff/drafts/{dr['id']}", json={"body": rebuilt, "edited": False}, headers=CSRF).json()
    assert up["edited"] is False and "December 1, 2026" in up["body"]
    up = staff_client.put(f"/api/staff/drafts/{dr['id']}", json={"status": "reviewed", "title": "NOV fence"}, headers=CSRF).json()
    assert up["status"] == "reviewed" and up["title"] == "NOV fence" and up["reviewed_by"] == "alice"

    # Linking to a second case adds one timeline item there.
    other = make_case(staff_client, address="3 Elm Street")
    staff_client.put(f"/api/staff/drafts/{dr['id']}", json={"case_id": other["id"]}, headers=CSRF)
    assert len(staff_client.get(f"/api/staff/cases/{other['id']}").json()["items"]) == 1

    assert staff_client.delete(f"/api/staff/drafts/{dr['id']}", headers=CSRF).json() == {"ok": True}
    assert staff_client.get(f"/api/staff/drafts/{dr['id']}").status_code == 404
    assert staff_client.get(f"/api/staff/cases/{case['id']}").json()["items"] == []
    assert staff_client.get(f"/api/staff/cases/{other['id']}").json()["items"] == []


def test_create_validation(staff_client):
    assert staff_client.post("/api/staff/drafts", json={"template": "nope"}, headers=CSRF).status_code == 404
    assert staff_client.post("/api/staff/drafts", json={"template": "nov-1", "case_id": "2026-000000"}, headers=CSRF).status_code == 404
    assert staff_client.post("/api/staff/drafts", json={"template": "nov-1", "values": {"bogus": "1"}}, headers=CSRF).status_code == 422
    assert staff_client.post("/api/staff/drafts", json={"template": "nov-1", "extra": 1}, headers=CSRF).status_code == 422
    assert staff_client.get("/api/staff/drafts/bad id!").status_code == 404


def test_docx_export_opens_with_letterhead_and_draft_footer(staff_client):
    dr = make_draft(staff_client, values=NOV_VALUES)
    r = staff_client.get(f"/api/staff/drafts/{dr['id']}/export.docx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml")
    assert f"DRAFT-nov-1-{dr['id']}.docx" in r.headers["content-disposition"]
    doc = Document(io.BytesIO(r.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Notice of violation and order to correct" in text
    assert "§ 275-6.2E(1)" in text and "A fence about 8 feet high" in text
    sec = doc.sections[0]
    header = "\n".join(p.text for p in sec.header.paragraphs)
    assert "CITY OF WATERVILLE" in header and "Code Enforcement Office" in header and "7 College Avenue" in header
    assert "207-680-4208" in header
    footer = "\n".join(p.text for p in sec.footer.paragraphs)
    assert footer.startswith("DRAFT FOR REVIEW.") and "Check every citation" in footer
    assert doc.core_properties.subject == "DRAFT for review"
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "Your right to appeal" in headings
    assert staff_client.get(f"/api/staff/drafts/{dr['id']}").json()["exported_by"] == "alice"


def test_docx_highlights_flags_and_renders_checklist(staff_client):
    from docx.enum.text import WD_COLOR_INDEX

    dr = make_draft(staff_client, template="80k-packet", values={"violation_type": "zoning_rp", "penalty_tier": "B-1"})
    doc = Document(io.BytesIO(staff_client.get(f"/api/staff/drafts/{dr['id']}/export.docx").content))
    runs = [r for p in doc.paragraphs for r in p.runs]
    assert any(r.text.startswith("[VERIFY") and r.font.highlight_color == WD_COLOR_INDEX.YELLOW for r in runs)
    assert any(r.text.startswith("[MISSING") and r.font.highlight_color == WD_COLOR_INDEX.PINK for r in runs)
    assert any(p.text.startswith("☐") for p in doc.paragraphs)
    header = "\n".join(p.text for p in doc.sections[0].header.paragraphs)
    assert "internal working document" in header  # no letterhead on the checklist
    assert doc.sections[0].footer.paragraphs[0].text.startswith("DRAFT")


def test_every_template_exports(staff_client):
    for tid in d.templates():
        dr = make_draft(staff_client, template=tid)
        r = staff_client.get(f"/api/staff/drafts/{dr['id']}/export.docx")
        assert r.status_code == 200, tid
        Document(io.BytesIO(r.content))


# ---------------------------------------------------------------- AI facts


def test_facts_fake_mode_uses_case_notes(staff_client):
    case = make_case(staff_client)
    staff_client.post(f"/api/staff/cases/{case['id']}/items", json={"kind": "note", "text": "Two unregistered cars by the barn."}, headers=CSRF)
    r = staff_client.post("/api/staff/drafts/facts", json={"template": "nov-1", "case_id": case["id"],
                                                         "values": {"inspection_date": "2026-09-30"}}, headers=CSRF)
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["field"] == "facts"
    assert "September 30, 2026" in j["text"] and "Two unregistered cars" in j["text"] and "[FACT NEEDED" in j["text"]


def test_facts_prompt_carries_case_record(staff_client, monkeypatch):
    seen = {}

    async def fake_complete(messages, max_tokens=1500, kind="complete"):
        seen.update(messages=messages, kind=kind)
        return "  On September 30, 2026, the Code Enforcement Officer observed a fence.  "

    monkeypatch.setattr(config, "FAKE_AZURE", False)
    monkeypatch.setattr(llm, "complete", fake_complete)
    case = make_case(staff_client)
    staff_client.post(f"/api/staff/cases/{case['id']}/items", json={"kind": "note", "text": "Fence 8 ft at rear line."}, headers=CSRF)
    r = staff_client.post("/api/staff/drafts/facts", json={"template": "nov-1", "case_id": case["id"],
                                                         "values": {"section_cited": "§ 275-4.26B(2)"}, "instructions": "Keep it short."}, headers=CSRF)
    assert r.status_code == 200, r.text
    assert r.json()["text"] == "On September 30, 2026, the Code Enforcement Officer observed a fence."
    assert seen["kind"] == "draft_facts"
    system, user = seen["messages"][0]["content"], seen["messages"][1]["content"]
    assert "Do not invent" in system and "[FACT NEEDED" in system and "data" in system
    assert "Fence 8 ft at rear line." in user and "12 Main Street" in user and "Keep it short." in user
    assert "§ 275-4.26B(2)" in user


def test_facts_errors(staff_client):
    assert staff_client.post("/api/staff/drafts/facts", json={"template": "zba-hearing"}, headers=CSRF).status_code == 422
    assert staff_client.post("/api/staff/drafts/facts", json={"template": "nov-1", "case_id": "2026-000000"}, headers=CSRF).status_code == 404
    assert staff_client.post("/api/staff/drafts/facts", json={"template": "nope"}, headers=CSRF).status_code == 404
